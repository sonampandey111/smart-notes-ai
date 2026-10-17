import os
import io
import json
import time
import uuid
import datetime
import urllib.request
import urllib.error
from flask import Flask, request, jsonify
from flask_cors import CORS
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

app = Flask(__name__)
# Enable CORS for all routes (Vite dev server and production)
CORS(app, resources={r"/*": {"origins": "*"}})

# Configuration
API_KEY = os.environ.get("GEMINI_API_KEY", "")
MONGODB_URI = os.environ.get("MONGODB_URI", "mongodb://localhost:27017")
DB_NAME = "smart_notes_ai"

# AI Models fallback order
MODELS_TO_TRY = [
    "models/gemini-flash-latest",
    "models/gemini-2.5-flash-lite",
    "models/gemma-4-26b-a4b-it",
    "models/gemma-4-31b-it"
]

# -------------------------------------------------------------
# DATABASE LAYER (MongoDB with Automatic Local JSON Fallback)
# -------------------------------------------------------------
DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
os.makedirs(DATA_DIR, exist_ok=True)
LOCAL_DB_FILE = os.path.join(DATA_DIR, "db.json")

def load_local_db():
    if not os.path.exists(LOCAL_DB_FILE):
        initial = {
            "notes": [],
            "stats": {
                "total_notes": 0,
                "total_words": 0,
                "quizzes_generated": 0,
                "flashcards_generated": 0
            }
        }
        with open(LOCAL_DB_FILE, "w", encoding="utf-8") as f:
            json.dump(initial, f, indent=2)
        return initial
    try:
        with open(LOCAL_DB_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {"notes": [], "stats": {"total_notes": 0, "total_words": 0, "quizzes_generated": 0, "flashcards_generated": 0}}

def save_local_db(data):
    with open(LOCAL_DB_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

mongo_client = None
mongo_db = None
mongo_available = False

try:
    from pymongo import MongoClient
    mongo_client = MongoClient(MONGODB_URI, serverSelectionTimeoutMS=2000)
    mongo_client.server_info()  # Test connection
    mongo_db = mongo_client[DB_NAME]
    mongo_available = True
    print(f"[Database] Connected to MongoDB at {MONGODB_URI} (Database: {DB_NAME})")
except Exception as e:
    mongo_available = False
    print(f"[Database] MongoDB not reachable ({e}). Using robust Local JSON Database fallback.")

def db_save_note(note_data):
    """Saves note record to MongoDB or Local DB."""
    if mongo_available and mongo_db is not None:
        try:
            doc_to_save = dict(note_data)
            # Use string ID for _id to ensure clean JSON serialization
            doc_to_save["_id"] = note_data["id"]
            mongo_db.notes.insert_one(doc_to_save)
            # Update stats in mongo
            mongo_db.stats.update_one(
                {"_id": "global_stats"},
                {
                    "$inc": {
                        "total_notes": 1,
                        "total_words": note_data.get("word_count", 0),
                        "quizzes_generated": len(note_data.get("quiz", {}).get("mcqs", [])) + len(note_data.get("quiz", {}).get("true_false", [])),
                        "flashcards_generated": len(note_data.get("flashcards", []))
                    }
                },
                upsert=True
            )
            note_data.pop("_id", None)
            return note_data["id"]
        except Exception as e:
            print(f"[DB Error] MongoDB insert failed: {e}. Falling back to local storage.")
            note_data.pop("_id", None)

    # Local DB fallback
    local_data = load_local_db()
    clean_note = {k: v for k, v in note_data.items() if k != "_id"}
    local_data["notes"].insert(0, clean_note)
    stats = local_data.setdefault("stats", {
        "total_notes": 0, "total_words": 0, "quizzes_generated": 0, "flashcards_generated": 0
    })
    stats["total_notes"] += 1
    stats["total_words"] += note_data.get("word_count", 0)
    stats["quizzes_generated"] += len(note_data.get("quiz", {}).get("mcqs", [])) + len(note_data.get("quiz", {}).get("true_false", []))
    stats["flashcards_generated"] += len(note_data.get("flashcards", []))
    save_local_db(local_data)
    return note_data["id"]

def db_get_all_notes(limit=25):
    """Retrieves list of notes summaries."""
    if mongo_available and mongo_db is not None:
        try:
            cursor = mongo_db.notes.find({}, {"_id": 0, "raw_text": 0}).sort("created_at", -1).limit(limit)
            return list(cursor)
        except Exception as e:
            print(f"[DB Error] MongoDB find failed: {e}")

    local_data = load_local_db()
    notes = []
    for n in local_data.get("notes", [])[:limit]:
        item = {k: v for k, v in n.items() if k != "raw_text"}
        notes.append(item)
    return notes

def db_get_note_by_id(note_id):
    """Retrieves full note by ID."""
    if mongo_available and mongo_db is not None:
        try:
            res = mongo_db.notes.find_one({"id": note_id}, {"_id": 0})
            if res:
                return res
        except Exception as e:
            print(f"[DB Error] MongoDB find_one failed: {e}")

    local_data = load_local_db()
    for n in local_data.get("notes", []):
        if n.get("id") == note_id:
            return n
    return None

def db_delete_note(note_id):
    """Deletes note by ID."""
    if mongo_available and mongo_db is not None:
        try:
            mongo_db.notes.delete_one({"id": note_id})
        except Exception as e:
            print(f"[DB Error] MongoDB delete failed: {e}")

    local_data = load_local_db()
    local_data["notes"] = [n for n in local_data.get("notes", []) if n.get("id") != note_id]
    save_local_db(local_data)
    return True

def db_get_stats():
    """Retrieves aggregated stats."""
    if mongo_available and mongo_db is not None:
        try:
            stats_doc = mongo_db.stats.find_one({"_id": "global_stats"})
            if stats_doc:
                return {
                    "total_notes": stats_doc.get("total_notes", 0),
                    "total_words": stats_doc.get("total_words", 0),
                    "quizzes_generated": stats_doc.get("quizzes_generated", 0),
                    "flashcards_generated": stats_doc.get("flashcards_generated", 0),
                    "db_type": "MongoDB (Active)"
                }
        except Exception as e:
            print(f"[DB Error] MongoDB stats failed: {e}")

    local_data = load_local_db()
    stats = local_data.get("stats", {
        "total_notes": 0, "total_words": 0, "quizzes_generated": 0, "flashcards_generated": 0
    })
    stats["db_type"] = "Local Storage (Active)"
    return stats

# -------------------------------------------------------------
# FILE PARSING FUNCTIONS (PDF, DOCX, TXT)
# -------------------------------------------------------------
def extract_text_from_pdf(file_bytes):
    """Extracts plain text from PDF using pypdf."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    extracted = []
    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            extracted.append(page_text.strip())
    if not extracted:
        raise ValueError("The PDF appears to be empty or contains only scanned images.")
    return "\n\n".join(extracted)

def extract_text_from_docx(file_bytes):
    """Extracts plain text from DOCX files using python-docx."""
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            extracted.append(t)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                extracted.append(" | ".join(row_text))
    if not extracted:
        raise ValueError("The DOCX file appears to have no readable text.")
    return "\n\n".join(extracted)

def extract_text_from_txt(file_bytes):
    """Decodes plain text bytes with fallback encodings."""
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")

# -------------------------------------------------------------
# AI ENGINE INTEGRATION (GEMINI / GEMMA)
# -------------------------------------------------------------
def call_ai(prompt, temperature=0.3, max_tokens=3000):
    """Calls Gemini/Gemma API with fallback support across model endpoints."""
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": temperature,
            "maxOutputTokens": max_tokens
        }
    }
    encoded_data = json.dumps(payload).encode("utf-8")
    last_error = None

    for model_name in MODELS_TO_TRY:
        url = f"https://generativelanguage.googleapis.com/v1beta/{model_name}:generateContent?key={API_KEY}"
        req = urllib.request.Request(
            url,
            data=encoded_data,
            headers={"Content-Type": "application/json"}
        )
        try:
            with urllib.request.urlopen(req, timeout=50) as response:
                res_data = json.loads(response.read().decode("utf-8"))
                candidates = res_data.get("candidates", [])
                if candidates and "content" in candidates[0]:
                    parts = candidates[0]["content"].get("parts", [])
                    if parts and "text" in parts[0]:
                        raw = parts[0]["text"].strip()
                        # Clean up internal thoughts if present
                        if "</thought>" in raw:
                            raw = raw.split("</thought>")[-1].strip()
                        return raw
        except urllib.error.HTTPError as http_err:
            body = http_err.read().decode("utf-8", errors="ignore")
            last_error = f"{model_name} (HTTP {http_err.code}): {body}"
            continue
        except Exception as e:
            last_error = f"{model_name}: {str(e)}"
            continue

    raise RuntimeError(f"All AI model requests failed. Last detail: {last_error}")

def extract_json_safely(text):
    """Strips markdown code fences and parses JSON safely."""
    clean = text.strip()
    if "</thought>" in clean:
        clean = clean.split("</thought>")[-1].strip()
    if "```json" in clean:
        clean = clean.split("```json", 1)[1].split("```", 1)[0].strip()
    elif "```" in clean:
        clean = clean.split("```", 1)[1].split("```", 1)[0].strip()
    
    start_idx = clean.find("{")
    end_idx = clean.rfind("}")
    if start_idx != -1 and end_idx != -1 and end_idx > start_idx:
        json_str = clean[start_idx:end_idx + 1]
        try:
            return json.loads(json_str)
        except Exception:
            pass
    return json.loads(clean)

# -------------------------------------------------------------
# API ROUTES
# -------------------------------------------------------------
@app.route("/", methods=["GET"])
@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({
        "status": "online",
        "service": "Smart AI Notes Analyzer API",
        "mongodb_connected": mongo_available,
        "models": MODELS_TO_TRY,
        "timestamp": datetime.datetime.now().isoformat()
    })

@app.route("/api/stats", methods=["GET"])
def get_stats():
    try:
        stats = db_get_stats()
        return jsonify({"success": True, "stats": stats})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/history", methods=["GET"])
def get_history():
    try:
        limit = request.args.get("limit", default=25, type=int)
        notes = db_get_all_notes(limit=limit)
        return jsonify({"success": True, "notes": notes})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/notes/<note_id>", methods=["GET"])
def get_single_note(note_id):
    try:
        note = db_get_note_by_id(note_id)
        if not note:
            return jsonify({"error": "Note not found"}), 404
        return jsonify({"success": True, "note": note})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/notes/<note_id>", methods=["DELETE"])
def delete_single_note(note_id):
    try:
        db_delete_note(note_id)
        return jsonify({"success": True, "message": "Note deleted successfully"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/analyze", methods=["POST"])
def analyze_notes():
    """
    Full Analysis Pipeline:
    1. Extracts text from PDF, DOCX, TXT or Textarea
    2. Generates:
       - AI Summarization (Overview, Core Concepts, Study Checklist)
       - Key Points (Concepts, Definitions, Formulas, Keywords)
       - AI Quiz (MCQs, True/False, Short Answer)
       - Flashcards (Interactive revision cards)
    3. Handles Language: English ('en') or Hindi ('hi')
    4. Saves to MongoDB / Local DB
    """
    try:
        raw_text = ""
        source_name = "Pasted Notes"

        # 1. Check direct text input
        form_text = request.form.get("text", "").strip()
        if form_text:
            raw_text += form_text

        # 2. Check uploaded file
        uploaded_file = request.files.get("file")
        if uploaded_file and uploaded_file.filename:
            filename = uploaded_file.filename
            source_name = filename
            lower_name = filename.lower()
            file_bytes = uploaded_file.read()

            if len(file_bytes) > 20 * 1024 * 1024:
                return jsonify({"error": "File size exceeds 20MB limit."}), 400

            if lower_name.endswith(".pdf"):
                pdf_text = extract_text_from_pdf(file_bytes)
                raw_text = f"{raw_text}\n\n{pdf_text}".strip()
            elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
                docx_text = extract_text_from_docx(file_bytes)
                raw_text = f"{raw_text}\n\n{docx_text}".strip()
            elif lower_name.endswith(".txt"):
                txt_text = extract_text_from_txt(file_bytes)
                raw_text = f"{raw_text}\n\n{txt_text}".strip()
            else:
                return jsonify({"error": "Unsupported file format. Please upload PDF, DOCX, or TXT."}), 400

        if not raw_text.strip():
            return jsonify({"error": "Please provide notes by pasting text or uploading a PDF, DOCX, or TXT file."}), 400

        # Language selection ('en' or 'hi')
        language = request.form.get("language", "en").strip().lower()
        lang_instruction = "Respond entirely in clear, natural Hindi (हिन्दी भाषा)." if language == "hi" else "Respond in clear English."

        word_count = len(raw_text.split())
        # Safe length truncation for API token limits
        content_for_ai = raw_text[:35000] if len(raw_text) > 35000 else raw_text

        # --- STEP 1: GENERATE STRUCTURED ANALYSIS (Summary & Key Points) ---
        prompt_analysis = f"""You are an elite academic tutor and study notes analyzer.
Language Requirement: {lang_instruction}

Analyze the following study material and return a JSON object with this exact structure:
{{
  "title": "A short, descriptive title for these notes (max 8 words)",
  "executive_summary": "A concise, high-impact 2-3 paragraph summary of the entire material explaining core themes clearly.",
  "key_concepts": [
    "Core concept 1 explained clearly with context",
    "Core concept 2 explained clearly with context",
    "Core concept 3 explained clearly with context",
    "Core concept 4 explained clearly with context"
  ],
  "definitions": [
    {{"term": "Term Name", "definition": "Clear academic definition"}},
    {{"term": "Term Name", "definition": "Clear academic definition"}}
  ],
  "formulas_and_laws": [
    {{"name": "Law / Equation Name", "formula": "Mathematical or Scientific notation", "description": "What it means and how it is applied"}}
  ],
  "keywords": ["Keyword1", "Keyword2", "Keyword3", "Keyword4", "Keyword5", "Keyword6"],
  "study_checklist": [
    "Actionable review question or task for students to verify understanding",
    "Another study checklist item"
  ]
}}

Ensure valid JSON only. Do not wrap in extra explanation.

Notes Content:
-----------------------------
{content_for_ai}
-----------------------------
"""

        analysis_raw = call_ai(prompt_analysis, temperature=0.25)
        try:
            analysis_data = extract_json_safely(analysis_raw)
        except Exception:
            # Fallback parsing
            analysis_data = {
                "title": source_name if source_name != "Pasted Notes" else "Study Notes Analysis",
                "executive_summary": analysis_raw,
                "key_concepts": ["Review comprehensive executive summary."],
                "definitions": [],
                "formulas_and_laws": [],
                "keywords": ["Study", "Revision", "Notes"],
                "study_checklist": ["Review the core summary concepts."]
            }

        # --- STEP 2: GENERATE INTERACTIVE QUIZ & FLASHCARDS ---
        prompt_quiz_cards = f"""You are a master educator creating high-retention revision materials.
Language Requirement: {lang_instruction}

Based on the study notes below, generate a JSON object containing:
1. "mcqs": Exactly 5 Multiple Choice Questions. Each with:
   - "question": The question string
   - "options": Array of exactly 4 choices [Choice A, Choice B, Choice C, Choice D]
   - "answerIndex": Integer (0, 1, 2, or 3) indicating the correct option
   - "explanation": Why this answer is correct

2. "true_false": Exactly 3 True/False questions. Each with:
   - "statement": The statement to evaluate
   - "answer": Boolean (true or false)
   - "explanation": Explanation of the truth value

3. "short_answer": Exactly 2 conceptual questions. Each with:
   - "question": Thought-provoking question
   - "sampleAnswer": Ideal model answer
   - "keyPoints": Array of 2-3 essential phrases/facts that must be mentioned

4. "flashcards": Exactly 6 to 8 flashcards for active recall. Each with:
   - "front": Question, concept, or term to test
   - "back": Clear, concise answer or definition (1-2 sentences)
   - "category": Topic tag (e.g. "Definition", "Mechanism", "Formula", "Fact")

JSON Structure:
{{
  "mcqs": [...],
  "true_false": [...],
  "short_answer": [...],
  "flashcards": [...]
}}

Output ONLY raw valid JSON.

Notes Content:
-----------------------------
{content_for_ai}
-----------------------------
"""

        quiz_cards_raw = call_ai(prompt_quiz_cards, temperature=0.2)
        try:
            quiz_cards_data = extract_json_safely(quiz_cards_raw)
        except Exception:
            quiz_cards_data = {
                "mcqs": [
                    {
                        "question": "What is the primary focus of these study notes?",
                        "options": ["Fundamental concepts", "Historical trivia", "Unrelated topics", "None of the above"],
                        "answerIndex": 0,
                        "explanation": "The material introduces core academic definitions and principles."
                    }
                ],
                "true_false": [
                    {
                        "statement": "Active recall and flashcard revision improve retention.",
                        "answer": True,
                        "explanation": "Scientific studies show active testing is more effective than passive reading."
                    }
                ],
                "short_answer": [
                    {
                        "question": "Summarize the primary thesis of the notes in your own words.",
                        "sampleAnswer": "The notes detail fundamental theories, definitions, and operational mechanisms.",
                        "keyPoints": ["Core mechanism", "Key definitions"]
                    }
                ],
                "flashcards": [
                    {
                        "front": "Main Concept",
                        "back": "Review the summary overview for the fundamental idea.",
                        "category": "Concept"
                    }
                ]
            }

        note_id = str(uuid.uuid4())
        created_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

        result_payload = {
            "id": note_id,
            "title": analysis_data.get("title", source_name),
            "source_type": source_name,
            "language": language,
            "word_count": word_count,
            "created_at": created_at,
            "summary": {
                "executive_summary": analysis_data.get("executive_summary", ""),
                "study_checklist": analysis_data.get("study_checklist", [])
            },
            "key_points": {
                "concepts": analysis_data.get("key_concepts", []),
                "definitions": analysis_data.get("definitions", []),
                "formulas": analysis_data.get("formulas_and_laws", []),
                "keywords": analysis_data.get("keywords", [])
            },
            "quiz": {
                "mcqs": quiz_cards_data.get("mcqs", []),
                "true_false": quiz_cards_data.get("true_false", []),
                "short_answer": quiz_cards_data.get("short_answer", [])
            },
            "flashcards": quiz_cards_data.get("flashcards", []),
            "raw_text": raw_text
        }

        # Save to database (MongoDB or local)
        db_save_note(result_payload)

        return jsonify({
            "success": True,
            "data": result_payload
        })

    except ValueError as ve:
        return jsonify({"error": str(ve)}), 400
    except RuntimeError as re:
        return jsonify({"error": str(re)}), 500
    except Exception as e:
        return jsonify({"error": f"An unexpected error occurred: {str(e)}"}), 500

@app.route("/api/ask", methods=["POST"])
def ask_question():
    """
    Q&A Chatbot: Answers student questions strictly grounded in the uploaded notes.
    """
    try:
        data = request.get_json(force=True)
        question = data.get("question", "").strip()
        note_id = data.get("note_id")
        context_text = data.get("context", "")
        language = data.get("language", "en")
        chat_history = data.get("history", [])

        if not question:
            return jsonify({"error": "Question is required."}), 400

        # If note_id is provided, retrieve the original text
        if note_id:
            note = db_get_note_by_id(note_id)
            if note and "raw_text" in note:
                context_text = note["raw_text"]

        if not context_text:
            return jsonify({"error": "No notes context found to answer the question."}), 400

        lang_instruction = "Answer in Hindi (हिन्दी)." if language == "hi" else "Answer in English."

        history_formatted = ""
        if chat_history:
            history_formatted = "Previous conversation:\n" + "\n".join(
                [f"{m.get('sender', 'user')}: {m.get('text', '')}" for m in chat_history[-4:]]
            ) + "\n\n"

        prompt = f"""You are a helpful and accurate academic study partner.
Language Requirement: {lang_instruction}

Answer the student's question based strictly on the provided study notes.
If the notes do not contain the answer, politely state that the answer is not mentioned in the uploaded notes, and provide a helpful general explanation.
Use formatting, bullet points, or bold text where appropriate for clarity.

Study Notes Context:
---------------------------
{context_text[:30000]}
---------------------------

{history_formatted}Student Question: {question}
Answer:"""

        answer_text = call_ai(prompt, temperature=0.3, max_tokens=1500)

        return jsonify({
            "success": True,
            "answer": answer_text
        })

    except Exception as e:
        return jsonify({"error": str(e)}), 500

# Legacy compatibility route for previous /summarize
@app.route("/summarize", methods=["POST"])
def legacy_summarize():
    return analyze_notes()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    print(f"============================================================")
    print(f"🚀 Smart AI Notes Analyzer Flask Server starting...")
    print(f"📡 URL: http://127.0.0.1:{port}")
    print(f"📦 MongoDB: {'Connected' if mongo_available else 'Fallback Local DB active'}")
    print(f"============================================================")
    app.run(host="127.0.0.1", port=port, debug=True)
